# -*- coding: utf-8 -*-
"""History Word document helpers with timeline support."""

import re
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx import Document


def setup_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(11)
    return doc


def _shade_cell(cell, color="D9E2F3"):
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn("w:shd"), {qn("w:fill"): color, qn("w:val"): "clear"})
    shading.append(shd)


def add_title_page(doc):
    title = doc.add_heading("人教版（统编版）初中历史 · 七～九年级中考总复习手册", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2022新课标 · 人教统编版（2024/2025/2026版）· 2026中考适用")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(89, 89, 89)
    doc.add_paragraph(
        "说明：初中历史为教育部统编版（封面：人民教育出版社）。"
        "七～八年级为中国史，九年级为世界史。"
        "七上2024秋、七下2025春、八上2024/2025秋、八下2025/2026春；"
        "九上2026秋新版、九下现行版。初二地生会考不考历史，历史在初三中考。"
    )
    doc.add_paragraph()


def add_book_title(doc, title):
    doc.add_page_break()
    doc.add_heading(title, 1)


def add_chapter_title(doc, title):
    doc.add_heading(title, 2)


def add_section_title(doc, code, title):
    p = doc.add_paragraph()
    r = p.add_run(f"{code} {title}")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(102, 0, 0)


def add_text_with_bold(doc, text, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Pt(18)
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part:
            p.add_run(part)


def add_plain(doc, text):
    add_text_with_bold(doc, text)


def add_fact(doc, example, concept):
    p = doc.add_paragraph()
    p.add_run(example)
    p.add_run(" → ")
    r = p.add_run(concept)
    r.bold = True


def add_block_title(doc, title, color=RGBColor(192, 0, 0)):
    p = doc.add_paragraph()
    r = p.add_run(f"【{title}】")
    r.bold = True
    r.font.color.rgb = color


def add_supplement(doc, items):
    add_block_title(doc, "补充说明", RGBColor(0, 112, 0))
    for item in items:
        add_text_with_bold(doc, item, indent=True)


def add_errors(doc, items):
    add_block_title(doc, "易错")
    for item in items:
        add_text_with_bold(doc, item, indent=True)


def add_exam_types(doc, items):
    add_block_title(doc, "常考题型", RGBColor(0, 112, 192))
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        _shade_cell(table.rows[0].cells[i])
        for r in table.rows[0].cells[i].paragraphs[0].runs:
            r.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = str(val)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def add_timeline_title(doc, title):
    p = doc.add_paragraph()
    r = p.add_run(f"▎{title}")
    r.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    r.font.size = Pt(11)


def add_timeline_arrow(doc, events):
    """Horizontal arrow timeline: [(time, event), ...]"""
    parts = []
    for i, (t, e) in enumerate(events):
        parts.append(f"{t} {e}")
    line = "  →  ".join(parts)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(12)
    r = p.add_run(line)
    r.font.size = Pt(10)
    doc.add_paragraph()


def add_timeline_table(doc, rows, title=None):
    if title:
        add_timeline_title(doc, title)
    add_table(doc, ["时间", "重大事件", "记忆要点/意义"], rows, [3.5, 5, 7])


def add_dynasty_song(doc):
    add_block_title(doc, "朝代歌（必背）", RGBColor(128, 0, 128))
    add_plain(doc, "夏商与西周，东周分两段；春秋和战国，一统秦两汉；")
    add_plain(doc, "三分魏蜀吴，二晋前后延；南北朝并立，隋唐五代传；")
    add_plain(doc, "宋元明清后，皇朝至此完。")
    doc.add_paragraph()
