# -*- coding: utf-8 -*-
"""初中数学 Word 文档 helpers."""

from docx.shared import Inches

from history_doc_helpers import *  # noqa: F401, F403


def add_title_page(doc):
    title = doc.add_heading("人教版初中数学 · 七～九年级总复习手册", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2022新课标 · 人教版 · 2026适用 · 按章节详解")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(89, 89, 89)
    doc.add_paragraph(
        "说明：本手册比课本**更详细**，按章节整理**概念、公式、方法、易错题、提升题**。"
        "孩子**马上初二**，请优先巩固**七年级全册**，再预习**八年级上册**。"
        "七上为2024秋新版目录；七下～九年级按现行人教版编排。"
        "含**几何模型、代数应用题拆题训练**专题。"
    )
    doc.add_paragraph()


def add_concept_block(doc, items):
    add_block_title(doc, "核心概念", RGBColor(0, 51, 102))
    for item in items:
        add_text_with_bold(doc, item, indent=True)


def add_formula_block(doc, formulas):
    """formulas: [(名称, 内容), ...]"""
    add_block_title(doc, "公式与定理", RGBColor(128, 0, 128))
    add_table(doc, ["名称", "内容"], formulas, [3, 10])


def add_method_block(doc, items):
    add_block_title(doc, "解题方法", RGBColor(0, 112, 0))
    for item in items:
        add_text_with_bold(doc, item, indent=True)


def add_mistake_problems(doc, problems):
    """易错题: [(题目, 错因, 正解), ...]"""
    add_block_title(doc, "易错题精讲", RGBColor(192, 0, 0))
    for i, (q, wrong, right) in enumerate(problems, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"【易错{i}】{q}")
        r.bold = True
        add_text_with_bold(doc, f"✗ 常见错法：{wrong}", indent=True)
        add_text_with_bold(doc, f"✓ 正确做法：{right}", indent=True)
        doc.add_paragraph()


def add_advanced_problems(doc, problems):
    """提升题: [(题目, 思路, 答案), ...]"""
    add_block_title(doc, "提升题（拔高练）", RGBColor(0, 112, 192))
    for i, (q, idea, ans) in enumerate(problems, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"【提升{i}】{q}")
        r.bold = True
        add_text_with_bold(doc, f"思路：{idea}", indent=True)
        add_text_with_bold(doc, f"答案：{ans}", indent=True)
        doc.add_paragraph()


def add_figure(doc, fig_path, caption=None, width_inches=4.8):
    """插入配图（居中）。"""
    if not fig_path:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(fig_path, width=Inches(width_inches))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(caption)
        cr.font.size = Pt(9)
        cr.font.color.rgb = RGBColor(89, 89, 89)


def add_model_card(doc, name, when, steps, example=None, figure_id=None):
    """几何/代数模型卡片。"""
    p = doc.add_paragraph()
    r = p.add_run(f"▎模型：{name}")
    r.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    if figure_id:
        from math_figure_draw import ensure_figure
        add_figure(doc, ensure_figure(figure_id), "示意图")
    add_text_with_bold(doc, f"**何时用**：{when}", indent=True)
    add_text_with_bold(doc, f"**步骤**：{steps}", indent=True)
    if example:
        add_text_with_bold(doc, f"**例题**：{example}", indent=True)
    doc.add_paragraph()


def add_decompose_drill(doc, num, problem, guide_steps, decompose_ans, solve_ans, figure_id=None):
    """拆题训练：配图 + 留白四步 + 参考答案。"""
    p = doc.add_paragraph()
    r = p.add_run(f"练习{num}　{problem}")
    r.bold = True
    if figure_id:
        from math_figure_draw import ensure_figure
        add_figure(doc, ensure_figure(figure_id), "图：先看图再拆题")
    add_block_title(doc, "我来拆题（先自己做，再对答案）", RGBColor(0, 112, 192))
    for step_name, hint in guide_steps:
        sp = doc.add_paragraph()
        sp.add_run(f"□ {step_name}（{hint}）")
        sp.paragraph_format.left_indent = Pt(18)
        doc.add_paragraph("_" * 40)
    add_block_title(doc, "参考答案·拆题思路", RGBColor(0, 112, 0))
    add_plain(doc, decompose_ans)
    add_block_title(doc, "参考答案·完整解答", RGBColor(128, 0, 128))
    add_plain(doc, solve_ans)
    doc.add_paragraph()
