# -*- coding: utf-8 -*-
"""初中物理 Word 文档 helpers."""

from docx.shared import Inches

from history_doc_helpers import *  # noqa: F401, F403


def add_title_page(doc):
    title = doc.add_heading("人教版初中物理 · 八～九年级中考总复习手册", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2022新课标 · 人教版 · 2026适用 · 按章节详解 + 拆题训练")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(89, 89, 89)
    doc.add_paragraph(
        "说明：物理从**初二上册**开始。本手册比课本**更详细**，每章含**概念、公式、实验、易错、常考题型**。"
        "孩子**马上初二**，请优先学**八上**（机械运动、声、物态、光、透镜、密度），再跟八下力学。"
        "含**物理拆题四步法**：圈词→判类型→找关系等式/隐藏数据→列式求解。"
    )
    doc.add_paragraph()


def add_concept_block(doc, items):
    add_block_title(doc, "核心概念", RGBColor(0, 51, 102))
    for item in items:
        add_text_with_bold(doc, item, indent=True)


def add_formula_block(doc, formulas):
    add_block_title(doc, "公式与规律", RGBColor(128, 0, 128))
    add_table(doc, ["名称", "内容"], formulas, [3, 10])


def add_method_block(doc, items):
    add_block_title(doc, "解题方法", RGBColor(0, 112, 0))
    for item in items:
        add_text_with_bold(doc, item, indent=True)


def add_experiment_block(doc, items):
    add_block_title(doc, "重点实验", RGBColor(0, 102, 102))
    for item in items:
        add_text_with_bold(doc, item, indent=True)


def add_mistake_problems(doc, problems):
    add_block_title(doc, "易错题精讲", RGBColor(192, 0, 0))
    for i, (q, wrong, right) in enumerate(problems, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"【易错{i}】{q}")
        r.bold = True
        add_text_with_bold(doc, f"✗ 常见错法：{wrong}", indent=True)
        add_text_with_bold(doc, f"✓ 正确做法：{right}", indent=True)
        doc.add_paragraph()


def add_advanced_problems(doc, problems):
    add_block_title(doc, "提升题（拔高练）", RGBColor(0, 112, 192))
    for i, (q, idea, ans) in enumerate(problems, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"【提升{i}】{q}")
        r.bold = True
        add_text_with_bold(doc, f"思路：{idea}", indent=True)
        add_text_with_bold(doc, f"答案：{ans}", indent=True)
        doc.add_paragraph()


def add_decompose_example(doc, problem, steps, relations, hidden, solution):
    """章节内拆题示范。"""
    p = doc.add_paragraph()
    r = p.add_run(f"▎拆题示范：{problem}")
    r.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_block_title(doc, "四步拆题", RGBColor(0, 112, 192))
    for name, content in steps:
        add_text_with_bold(doc, f"**{name}**：{content}", indent=True)
    add_block_title(doc, "关系等式", RGBColor(0, 112, 0))
    add_plain(doc, relations)
    add_block_title(doc, "隐藏数据/条件", RGBColor(192, 128, 0))
    add_plain(doc, hidden)
    add_block_title(doc, "完整解答", RGBColor(128, 0, 128))
    add_plain(doc, solution)
    doc.add_paragraph()


def add_figure(doc, fig_path, caption=None, width_inches=4.8):
    if not fig_path:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(fig_path, width=Inches(width_inches))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(caption)
        cr.font.size = Pt(9)
        cr.font.color.rgb = RGBColor(89, 89, 89)


def add_decompose_drill(doc, num, problem, guide_steps, decompose_ans, solve_ans, figure_id=None):
    """拆题专练：配图 + 留白四步 + 参考答案。"""
    p = doc.add_paragraph()
    r = p.add_run(f"练习{num}　{problem}")
    r.bold = True
    if figure_id:
        from physics_figure_draw import ensure_figure
        add_figure(doc, ensure_figure(figure_id), "图：先看图再拆题")
    add_block_title(doc, "我来拆题（先自己做，再对答案）", RGBColor(0, 112, 192))
    for step_name, hint in guide_steps:
        sp = doc.add_paragraph()
        sp.add_run(f"□ {step_name}（{hint}）")
        sp.paragraph_format.left_indent = Pt(18)
        doc.add_paragraph("_" * 40)
    add_block_title(doc, "参考答案·拆题思路", RGBColor(0, 112, 0))
    add_plain(doc, decompose_ans)
    add_block_title(doc, "参考答案·完整解答", RGBColor(128, 0, 128))
    add_plain(doc, solve_ans)
    doc.add_paragraph()
